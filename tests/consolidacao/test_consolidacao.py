import tempfile
import unittest
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

from app_captacao import consolidacao
from app_captacao.consolidacao.layouts import LAYOUTS
from app_captacao.consolidacao.leitura import ArquivoInvalido, ler_relatorio
from tests.consolidacao import planilhas_sienge as sienge

HOJE = date(2026, 10, 16)


def _abas(caminho) -> dict[str, list[tuple]]:
    livro = load_workbook(caminho, read_only=True)
    try:
        return {aba.title: list(aba.iter_rows(values_only=True)) for aba in livro.worksheets}
    finally:
        livro.close()  # no Windows, planilha aberta trava a pasta temporaria


def _dados(caminho) -> list[dict]:
    colunas, *linhas = _abas(caminho)["Dados"]
    return [dict(zip(colunas, linha)) for linha in linhas]


class _ComPasta(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.pasta = Path(self._tmp.name)

    def tearDown(self):
        self._tmp.cleanup()


class LeituraDoLayoutTests(_ComPasta):
    def test_apropriacoes_marca_entrada_e_saida_pela_posicao_do_bloco(self):
        lido = ler_relatorio(sienge.analitico(self.pasta / "a.xlsx"), LAYOUTS["analitico_propriedade"])
        resumo = [(linha["Conta"], linha["Natureza"], linha["Valor"]) for linha in lido.linhas]
        self.assertEqual(
            resumo,
            [
                ("1.01 - Receita de vendas", "Entrada", 100.0),
                ("1.01 - Receita de vendas", "Entrada", 50.25),
                ("2.12 - Tarifas bancárias", "Saída", 1.0),
            ],
        )
        self.assertEqual(lido.linhas[0]["Centro de custo"], "1002 - OBRA TESTE")
        self.assertEqual(lido.linhas[0]["Data"], date(2026, 8, 3))

    def test_extrato_baixa_adicional_herda_a_parcela_sem_repetir_o_valor_original(self):
        lido = ler_relatorio(sienge.extrato(self.pasta / "e.xlsx"), LAYOUTS["extrato_cliente_historico"])
        adicional = [linha for linha in lido.linhas if linha["Linha"] == "Baixa adicional"]
        self.assertEqual(len(adicional), 1)
        self.assertEqual((adicional[0]["Par"], adicional[0]["Tipo Condição"]), (2, "PARCELAS MENSAIS"))
        self.assertIsNone(adicional[0]["Valor original"])
        self.assertEqual(adicional[0]["Valor baixa"], 200.0)

    def test_extrato_texto_do_sienge_na_coluna_de_valor_vai_para_observacao(self):
        lido = ler_relatorio(sienge.extrato(self.pasta / "e.xlsx"), LAYOUTS["extrato_cliente_historico"])
        reparcelada = next(linha for linha in lido.linhas if linha["Título"] == 1 and linha["Par"] == 3)
        self.assertEqual(reparcelada["Observação"], "Reparcelamento")
        self.assertIsNone(reparcelada["Acréscimo"])

    def test_extrato_campos_de_um_titulo_nao_vazam_para_o_seguinte(self):
        lido = ler_relatorio(sienge.extrato(self.pasta / "e.xlsx"), LAYOUTS["extrato_cliente_historico"])
        por_titulo = {linha["Título"]: linha for linha in lido.linhas}
        self.assertEqual(por_titulo[1]["Corretor"], "CORRETORA X")
        self.assertIsNone(por_titulo[2]["Corretor"])  # o bloco do titulo 2 nao tem corretor
        self.assertEqual(por_titulo[2]["Cliente(s)"], "CLIENTE B(2)")

    def test_titulo_filtros_totais_e_rodape_nao_viram_dados(self):
        lido = ler_relatorio(sienge.unidades(self.pasta / "u.xlsx"), LAYOUTS["unidades"])
        self.assertEqual([linha["Unidade"] for linha in lido.linhas], ["Apto 101", "Apto 102", "Apto 103"])

    def test_arquivo_sem_o_cabecalho_esperado_e_recusado(self):
        caminho = sienge.unidades(self.pasta / "u.xlsx")
        with self.assertRaises(ArquivoInvalido) as erro:
            ler_relatorio(caminho, LAYOUTS["preco_total_vendas"])
        self.assertIn("u.xlsx", str(erro.exception))

    def test_arquivo_que_nao_e_planilha_e_recusado(self):
        caminho = self.pasta / "corrompido.xlsx"
        caminho.write_bytes(b"isto nao e um xlsx")
        with self.assertRaises(ArquivoInvalido):
            ler_relatorio(caminho, LAYOUTS["unidades"])


class ConferenciaTests(_ComPasta):
    def test_totais_que_batem_geram_o_consolidado(self):
        arquivos = [sienge.analitico(self.pasta / "a1.xlsx"), sienge.analitico(self.pasta / "a2.xlsx", centro="1003 - OBRA B")]
        resultado = consolidacao.consolidar("analitico_propriedade", "Analítico de Apropriações", arquivos, HOJE)
        self.assertTrue(resultado.gerado, resultado.motivo)
        self.assertEqual((resultado.linhas, resultado.arquivos), (6, 2))
        self.assertTrue(all(item.ok for item in resultado.conferencias))

    def test_total_impresso_diferente_da_soma_nao_gera_consolidado(self):
        arquivos = [sienge.analitico(self.pasta / "a.xlsx", total_saidas=11.0)]
        resultado = consolidacao.consolidar("analitico_propriedade", "Analítico de Apropriações", arquivos, HOJE)
        self.assertFalse(resultado.gerado)
        self.assertIn("Saídas somou 1, o Sienge imprimiu 11", resultado.motivo)
        self.assertEqual(list(self.pasta.glob("*consolidado*")), [])

    def test_diferenca_dentro_do_arredondamento_e_aceita(self):
        # 2 linhas: tolerancia de R$ 0,01; o Sienge soma antes de arredondar cada valor.
        arquivos = [sienge.preco_total(self.pasta / "p.xlsx", valores=(100.0, 200.0), total_contrato=300.01)]
        resultado = consolidacao.consolidar("preco_total_vendas", "Preço Total", arquivos, HOJE)
        self.assertTrue(resultado.gerado, resultado.motivo)

    def test_diferenca_acima_do_arredondamento_e_recusada(self):
        arquivos = [sienge.preco_total(self.pasta / "p.xlsx", valores=(100.0, 200.0), total_contrato=300.05)]
        resultado = consolidacao.consolidar("preco_total_vendas", "Preço Total", arquivos, HOJE)
        self.assertFalse(resultado.gerado)
        self.assertIn("Vl. contrato", resultado.motivo)

    def test_contagem_de_unidades_precisa_bater_exatamente(self):
        arquivos = [sienge.unidades(self.pasta / "u.xlsx", total_unidades=4)]
        resultado = consolidacao.consolidar("unidades", "Unidades", arquivos, HOJE)
        self.assertFalse(resultado.gerado)
        self.assertIn("Quantidade de unidades somou 3, o Sienge imprimiu 4", resultado.motivo)

    def test_extrato_confere_os_totais_de_cada_titulo(self):
        ok = consolidacao.consolidar("extrato_cliente_historico", "Extrato", [sienge.extrato(self.pasta / "e.xlsx")], HOJE)
        self.assertTrue(ok.gerado, ok.motivo)
        divergente = sienge.extrato(self.pasta / "e2.xlsx", total_recto_titulo_1=1600.0)
        falha = consolidacao.consolidar("extrato_cliente_historico", "Extrato", [divergente], HOJE)
        self.assertFalse(falha.gerado)
        self.assertIn("Recto líquido", falha.motivo)

    def test_relatorio_sem_total_impresso_e_gerado_sem_conferencia(self):
        resultado = consolidacao.consolidar("mapa_venda", "Mapa de Vendas", [sienge.mapa_vendas(self.pasta / "m.xlsx")], HOJE)
        self.assertTrue(resultado.gerado)
        self.assertFalse(resultado.conferido)

    def test_um_arquivo_invalido_impede_o_consolidado_inteiro(self):
        bom = sienge.unidades(self.pasta / "u.xlsx")
        ruim = self.pasta / "ruim.xlsx"
        ruim.write_bytes(b"x")
        resultado = consolidacao.consolidar("unidades", "Unidades", [bom, ruim], HOJE)
        self.assertFalse(resultado.gerado)
        self.assertIn("ruim.xlsx", resultado.motivo)

    def test_relatorio_sem_layout_nao_e_suportado(self):
        self.assertFalse(consolidacao.suporta("extrato_conciliado"))
        self.assertFalse(consolidacao.consolidar("extrato_conciliado", "X", [self.pasta / "a.xlsx"]).gerado)


class PlanilhaGeradaTests(_ComPasta):
    def test_consolidado_tem_dados_origem_e_conferencia_na_pasta_do_relatorio(self):
        arquivos = [
            sienge.unidades(self.pasta / "Unidades_242_2026-10-16.xlsx"),
            sienge.unidades(self.pasta / "Unidades_242_2026-10-16_3.xlsx", empreendimento="1003 - OBRA B"),
        ]
        resultado = consolidacao.consolidar("unidades", "Unidades", arquivos, HOJE)
        caminho = Path(resultado.caminho)
        self.assertEqual(caminho, self.pasta / "Unidades_consolidado_2026-10-16.xlsx")

        abas = _abas(caminho)
        self.assertEqual(list(abas), ["Dados", "Origem", "Conferência"])
        dados = _dados(caminho)
        self.assertEqual(len(dados), 6)
        self.assertEqual({linha["Empreendimento"] for linha in dados}, {"1002 - OBRA TESTE", "1003 - OBRA B"})
        self.assertEqual(dados[0]["Arquivo de origem"], "Unidades_242_2026-10-16.xlsx")
        self.assertEqual([linha[2:] for linha in abas["Origem"][1:]], [("1002 - OBRA TESTE", 3), ("1003 - OBRA B", 3)])
        self.assertEqual([linha[-1] for linha in abas["Conferência"][1:]], ["OK", "OK"])

    def test_datas_viram_data_do_excel(self):
        resultado = consolidacao.consolidar("unidades", "Unidades", [sienge.unidades(self.pasta / "u.xlsx")], HOJE)
        livro = load_workbook(resultado.caminho)
        celula = livro["Dados"]["D2"]  # "Posição em"
        livro.close()
        self.assertEqual(celula.value.date(), date(2026, 10, 16))
        self.assertEqual(celula.number_format, "dd/mm/yyyy")

    def test_consolidado_existente_nao_e_sobrescrito(self):
        arquivo = sienge.unidades(self.pasta / "u.xlsx")
        primeiro = consolidacao.consolidar("unidades", "Unidades", [arquivo], HOJE)
        segundo = consolidacao.consolidar("unidades", "Unidades", [arquivo], HOJE)
        self.assertNotEqual(primeiro.caminho, segundo.caminho)
        self.assertTrue(segundo.caminho.endswith("Unidades_consolidado_2026-10-16_2.xlsx"))
        self.assertTrue(Path(primeiro.caminho).exists())


if __name__ == "__main__":
    unittest.main()
