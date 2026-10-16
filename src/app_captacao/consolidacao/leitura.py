"""Leitura de um relatorio exportado pelo Sienge no layout de impressao.

O arquivo tem titulo, um bloco de filtros ("Empresa | ... | 242 - ..."), a linha de
cabecalho, os dados, subtotais/totais e rodape. Aqui os dados viram linhas planas, com os
campos do bloco de filtros como colunas, e os totais ficam guardados para a conferencia.
"""

from __future__ import annotations

import re
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from app_captacao.consolidacao.layouts import Layout

COLUNA_ARQUIVO = "Arquivo de origem"
COLUNA_NATUREZA = "Natureza"
COLUNA_LINHA = "Linha"
COLUNA_OBSERVACAO = "Observação"

LINHA_PARCELA = "Parcela"
LINHA_BAIXA_ADICIONAL = "Baixa adicional"

_DATA = re.compile(r"^\d{2}/\d{2}/\d{4}$")
_RODAPE = re.compile(r"^\d{2}/\d{2}/\d{4} - \d{2}:\d{2}(:\d{2})?$")  # "16/10/2026 - 09:11:55 | SIENGE / STARIAN"
_PREFIXOS_DE_TOTAL = ("Total", "Diferença", "**")
_NATUREZA_POR_TOTAL = {"Total de entradas": "Entrada", "Total de saídas": "Saída"}


class ArquivoInvalido(ValueError):
    """O arquivo nao tem o layout esperado: o consolidado nao e gerado."""


@dataclass(frozen=True)
class LinhaTotal:
    rotulo: str
    secao: str
    valores: tuple


@dataclass
class ArquivoLido:
    caminho: Path
    colunas: list[str]
    linhas: list[dict]
    totais: list[LinhaTotal] = field(default_factory=list)

    def campo_do_bloco(self, *rotulos: str) -> str:
        """Primeiro valor preenchido dos rotulos pedidos (empresa, empreendimento...)."""
        for linha in self.linhas:
            for rotulo in rotulos:
                valor = linha.get(rotulo)
                if valor not in (None, ""):
                    return str(valor)
        return ""


def texto(valor) -> str:
    return valor.strip() if isinstance(valor, str) else ""


def numero(valor) -> float | None:
    if isinstance(valor, bool):
        return None
    if isinstance(valor, (int, float)):
        return float(valor)
    if isinstance(valor, str):
        try:
            return float(valor.strip())
        except ValueError:
            return None
    return None


def _valor(valor):
    if isinstance(valor, str):
        limpo = valor.strip()
        if _DATA.match(limpo):
            try:
                return datetime.strptime(limpo, "%d/%m/%Y").date()
            except ValueError:
                return limpo
        return limpo
    if isinstance(valor, datetime):
        return valor.date() if valor.time() == datetime.min.time() else valor
    return valor


def _vazia(linha) -> bool:
    return all(v is None or (isinstance(v, str) and not v.strip()) for v in linha)


def _eh_cabecalho_de_secao(linha) -> bool:
    """'Total Geral | Valor | Quantidade': abre um bloco de totais, sem numeros."""
    return all(isinstance(v, str) for v in linha[1:] if v not in (None, ""))


def _ler_planilha(caminho: Path) -> list[list]:
    try:
        planilha = load_workbook(caminho, read_only=True, data_only=True)
    except (InvalidFileException, zipfile.BadZipFile, KeyError, OSError) as exc:
        raise ArquivoInvalido(f"{caminho.name} nao abre como planilha do Excel") from exc
    try:
        return [list(linha) for linha in planilha.worksheets[0].iter_rows(values_only=True)]
    finally:
        planilha.close()


def _ler_bloco_de_filtros(linha, layout: Layout, contexto: dict) -> None:
    for indice, celula in enumerate(linha):
        rotulo = texto(celula)
        if rotulo not in layout.rotulos:
            continue
        if rotulo == layout.inicio_de_bloco:
            contexto.clear()
        seguintes = [v for v in linha[indice + 1:] if v not in (None, "")]
        valor = seguintes[0] if seguintes else None
        contexto[rotulo] = None if texto(valor) in layout.rotulos else _valor(valor)


def colunas_do_layout(layout: Layout, colunas_tabela: list[str]) -> list[str]:
    colunas = [COLUNA_ARQUIVO, *layout.rotulos]
    if layout.natureza_por_secao:
        colunas.append(COLUNA_NATUREZA)
    colunas += [c for c in colunas_tabela if c not in colunas]
    if layout.baixa_adicional:
        colunas += [COLUNA_LINHA, COLUNA_OBSERVACAO]
    return colunas


def ler_relatorio(caminho, layout: Layout) -> ArquivoLido:
    caminho = Path(caminho)
    contexto: dict = {}
    cabecalho: list[tuple[int, str]] = []
    em_tabela = False
    secao = ""
    linhas: list[dict] = []
    totais: list[LinhaTotal] = []
    sem_natureza: list[dict] = []
    parcela_anterior: dict | None = None

    for linha in _ler_planilha(caminho):
        primeiro = texto(linha[0]) if linha else ""

        if layout.natureza_por_secao and primeiro in _NATUREZA_POR_TOTAL:
            for registro in sem_natureza:
                registro[COLUNA_NATUREZA] = _NATUREZA_POR_TOTAL[primeiro]
            sem_natureza = []

        if primeiro == layout.primeiro_campo:
            cabecalho = [(j, texto(v)) for j, v in enumerate(linha) if texto(v)]
            em_tabela = True
            continue
        if _vazia(linha) or _RODAPE.match(primeiro):
            em_tabela = False
            continue
        if primeiro.startswith(_PREFIXOS_DE_TOTAL) or primeiro in layout.linhas_de_total:
            if _eh_cabecalho_de_secao(linha):
                secao = primeiro
            else:
                totais.append(LinhaTotal(primeiro, secao, tuple(linha)))
            continue
        if not em_tabela:
            _ler_bloco_de_filtros(linha, layout, contexto)
            continue

        registro: dict = {COLUNA_ARQUIVO: caminho.name}
        registro.update({rotulo: contexto.get(rotulo) for rotulo in layout.rotulos})
        observacao = ""
        for j, nome in cabecalho:
            valor = linha[j] if j < len(linha) else None
            if texto(valor).startswith("***"):
                # O Sienge escreve "*** Reparcelamento ***" dentro de uma coluna de valor.
                observacao = texto(valor).strip("*").strip()
                valor = None
            registro[nome] = _valor(valor)
        if not observacao:
            observacao = next((texto(v).strip("*").strip() for v in linha if texto(v).startswith("***")), "")
        if layout.natureza_por_secao:
            registro[COLUNA_NATUREZA] = ""
            sem_natureza.append(registro)
        if layout.baixa_adicional:
            if registro.get(layout.primeiro_campo) in (None, "") and parcela_anterior is not None:
                registro[COLUNA_LINHA] = LINHA_BAIXA_ADICIONAL
                registro[layout.primeiro_campo] = parcela_anterior.get(layout.primeiro_campo)
                if len(cabecalho) > 1:
                    nome_tipo = cabecalho[1][1]
                    registro[nome_tipo] = parcela_anterior.get(nome_tipo)
            else:
                registro[COLUNA_LINHA] = LINHA_PARCELA
                parcela_anterior = registro
            registro[COLUNA_OBSERVACAO] = observacao
        linhas.append(registro)

    if not cabecalho:
        raise ArquivoInvalido(f"{caminho.name} nao tem a linha de cabecalho '{layout.primeiro_campo}'")
    return ArquivoLido(caminho, colunas_do_layout(layout, [nome for _, nome in cabecalho]), linhas, totais)


def eh_data(valor) -> bool:
    return isinstance(valor, (date, datetime))
