"""Junta os arquivos de um relatorio numa planilha so, depois de conferir cada um com o Sienge.

Se uma conferencia falha, nada e gravado: e melhor faltar o consolidado do que ele trazer
numeros diferentes dos que o Sienge imprimiu.
"""

from __future__ import annotations

import logging
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from app_captacao.consolidacao.layouts import LAYOUTS, Conferencia, Layout
from app_captacao.consolidacao.leitura import (
    COLUNA_NATUREZA,
    ArquivoInvalido,
    ArquivoLido,
    eh_data,
    ler_relatorio,
    numero,
)
from app_captacao.core.shared.arquivos import montar_caminho_unico, sanitizar_nome_arquivo

logger = logging.getLogger(__name__)

# Cada valor exportado vem arredondado em 2 casas; o total do Sienge e calculado antes disso.
ARREDONDAMENTO_POR_LINHA = 0.005
TOLERANCIA_MINIMA = 0.01


@dataclass(frozen=True)
class ResultadoConferencia:
    arquivo: str
    descricao: str
    extraido: float
    impresso: float | None
    ok: bool
    motivo: str = ""


@dataclass
class ResultadoConsolidacao:
    gerado: bool
    caminho: str = ""
    linhas: int = 0
    arquivos: int = 0
    motivo: str = ""
    conferencias: list[ResultadoConferencia] = field(default_factory=list)

    @property
    def conferido(self) -> bool:
        return bool(self.conferencias)


def suporta(executor: str) -> bool:
    return executor in LAYOUTS


def formatar_valor(valor: float | None) -> str:
    if valor is None:
        return "-"
    if float(valor).is_integer() and abs(valor) < 1e12:
        return f"{int(valor):,}".replace(",", ".")
    return f"{valor:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def _conferir_uma(lido: ArquivoLido, conferencia: Conferencia) -> ResultadoConferencia:
    linhas = [
        linha for linha in lido.linhas if not conferencia.natureza or linha.get(COLUNA_NATUREZA) == conferencia.natureza
    ]
    if conferencia.campo is None:
        extraido, quantidade, tolerancia = float(len(linhas)), len(linhas), 0.0
    else:
        valores = [v for v in (numero(linha.get(conferencia.campo)) for linha in linhas) if v is not None]
        extraido, quantidade = round(sum(valores), 2), len(valores)
        tolerancia = max(TOLERANCIA_MINIMA, ARREDONDAMENTO_POR_LINHA * quantidade)

    totais = [
        total
        for total in lido.totais
        if total.rotulo in conferencia.rotulos_total and (not conferencia.secao or total.secao == conferencia.secao)
    ]
    nome_total = " + ".join(conferencia.rotulos_total)
    if conferencia.secao:
        nome_total = f"{nome_total} ({conferencia.secao})"
    if not totais:
        # Periodo sem movimento: sem linhas e sem total e o esperado.
        ok = quantidade == 0
        motivo = "" if ok else f"total '{nome_total}' nao encontrado"
        return ResultadoConferencia(lido.caminho.name, conferencia.descricao, extraido, None, ok, motivo)

    impressos = [
        numero(total.valores[conferencia.coluna_total]) if conferencia.coluna_total < len(total.valores) else None
        for total in totais
    ]
    if any(valor is None for valor in impressos):
        motivo = f"total '{nome_total}' sem numero na coluna esperada (o layout do Sienge pode ter mudado)"
        return ResultadoConferencia(lido.caminho.name, conferencia.descricao, extraido, None, False, motivo)

    impresso = round(sum(v for v in impressos if v is not None), 2)
    ok = abs(extraido - impresso) <= tolerancia + 1e-9
    motivo = "" if ok else (
        f"{conferencia.descricao} somou {formatar_valor(extraido)}, o Sienge imprimiu {formatar_valor(impresso)}"
    )
    return ResultadoConferencia(lido.caminho.name, conferencia.descricao, extraido, impresso, ok, motivo)


def conferir(lido: ArquivoLido, layout: Layout) -> list[ResultadoConferencia]:
    resultados = []
    if layout.natureza_por_secao:
        sem_natureza = sum(1 for linha in lido.linhas if not linha.get(COLUNA_NATUREZA))
        if sem_natureza:
            resultados.append(
                ResultadoConferencia(
                    lido.caminho.name,
                    "Natureza",
                    float(sem_natureza),
                    None,
                    False,
                    f"{sem_natureza} linha(s) fora de um bloco de entradas ou de saidas",
                )
            )
    resultados += [_conferir_uma(lido, conferencia) for conferencia in layout.conferencias]
    return resultados


def _celula(planilha, valor):
    if eh_data(valor):
        celula = WriteOnlyCell(planilha, value=valor)
        celula.number_format = "dd/mm/yyyy"
        return celula
    return valor


def _aba(livro, titulo: str, colunas: list[str], larguras: dict[int, float] | None = None):
    aba = livro.create_sheet(titulo)
    aba.freeze_panes = "A2"
    for indice, coluna in enumerate(colunas, start=1):
        largura = (larguras or {}).get(indice) or min(max(len(coluna) + 4, 12), 45)
        aba.column_dimensions[get_column_letter(indice)].width = largura
    cabecalho = []
    for coluna in colunas:
        celula = WriteOnlyCell(aba, value=coluna)
        celula.font = Font(bold=True)
        cabecalho.append(celula)
    aba.append(cabecalho)
    return aba


def _gravar(destino: Path, colunas: list[str], lidos: list[ArquivoLido], conferencias: list[ResultadoConferencia]):
    livro = Workbook(write_only=True)
    total_linhas = sum(len(lido.linhas) for lido in lidos)

    dados = _aba(livro, "Dados", colunas)
    dados.auto_filter.ref = f"A1:{get_column_letter(len(colunas))}{total_linhas + 1}"
    for lido in lidos:
        for linha in lido.linhas:
            dados.append([_celula(dados, linha.get(coluna)) for coluna in colunas])

    origem = _aba(livro, "Origem", ["Arquivo", "Empresa", "Empreendimento / Centro de custo", "Linhas de dados"], {1: 55, 2: 45, 3: 50})
    for lido in lidos:
        origem.append(
            [
                lido.caminho.name,
                lido.campo_do_bloco("Empresa"),
                lido.campo_do_bloco("Empreendimento", "Centro de custo"),
                len(lido.linhas),
            ]
        )

    aba_conferencia = _aba(
        livro, "Conferência", ["Arquivo", "Conferência", "Extraído", "Impresso pelo Sienge", "Situação"], {1: 55, 2: 30}
    )
    for item in conferencias:
        situacao = "OK" if item.ok else item.motivo
        if item.ok and item.impresso is None:
            situacao = "OK (sem movimento)"
        aba_conferencia.append([item.arquivo, item.descricao, item.extraido, item.impresso, situacao])
    if not conferencias:
        aba_conferencia.append(["", "O Sienge nao imprime total neste relatorio", None, None, "Sem conferencia"])

    livro.save(destino)


def consolidar(executor: str, nome_relatorio: str, arquivos, data: date | None = None) -> ResultadoConsolidacao:
    """Le, confere e grava o consolidado na pasta do primeiro arquivo."""
    layout = LAYOUTS.get(executor)
    caminhos = [Path(arquivo) for arquivo in arquivos if str(arquivo or "").strip()]
    if layout is None:
        return ResultadoConsolidacao(False, motivo="relatorio sem layout de consolidacao")
    if not caminhos:
        return ResultadoConsolidacao(False, motivo="nenhum arquivo para consolidar")

    try:
        lidos = [ler_relatorio(caminho, layout) for caminho in caminhos]
    except ArquivoInvalido as exc:
        return ResultadoConsolidacao(False, arquivos=len(caminhos), motivo=str(exc))

    conferencias = [item for lido in lidos for item in conferir(lido, layout)]
    falhas = [item for item in conferencias if not item.ok]
    if falhas:
        primeira = falhas[0]
        motivo = f"{primeira.arquivo}: {primeira.motivo}"
        if len(falhas) > 1:
            motivo += f" (e mais {len(falhas) - 1} divergencia(s))"
        logger.warning("Consolidado de '%s' nao gerado: %s", nome_relatorio, [f.motivo for f in falhas])
        return ResultadoConsolidacao(False, arquivos=len(caminhos), motivo=motivo, conferencias=conferencias)

    colunas = list(dict.fromkeys(coluna for lido in lidos for coluna in lido.colunas))
    nome_base = f"{sanitizar_nome_arquivo(nome_relatorio)}_consolidado_{(data or date.today()):%Y-%m-%d}"
    destino = montar_caminho_unico(caminhos[0].parent, nome_base, ".xlsx")
    try:
        _gravar(destino, colunas, lidos, conferencias)
    except Exception:
        destino.unlink(missing_ok=True)
        raise

    linhas = sum(len(lido.linhas) for lido in lidos)
    logger.info("Consolidado de '%s' gerado: %s (%s linhas de %s arquivos)", nome_relatorio, destino, linhas, len(lidos))
    return ResultadoConsolidacao(True, str(destino), linhas, len(lidos), conferencias=conferencias)
