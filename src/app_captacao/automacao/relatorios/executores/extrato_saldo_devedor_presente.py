import math
import re
import time
from csv import reader as csv_reader
from pathlib import Path

from openpyxl import Workbook, load_workbook
from selenium.common.exceptions import NoSuchElementException, TimeoutException
from selenium.webdriver.common.keys import Keys
from selenium.webdriver.common.by import By
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.support.ui import WebDriverWait


TAMANHO_PAGINA = "5000"
TIMEOUT_CARREGAMENTO_LENTO = 480
PAUSA_ESTABILIZACAO = 0.5
PAUSA_INICIO_DOWNLOAD = 2
TIMEOUT_LOCALIZACAO_CAMPO = 2
CAMPO_EMPRESA = "cdEmpresaView"
IDENTIFICADORES_EMPRESA = ("cdEmpresaView", "filter.empresa.cdEmpresaView")
CAMPO_DATA_INICIO = "dtInicio"
CAMPO_DATA_FIM = "dtFim"
IDENTIFICADORES_DATA_INICIO = ("dataInicio", "dtInicio", "filter.dtInicioStr")
IDENTIFICADORES_DATA_FIM = ("dataFim", "dtFim", "filter.dtFimStr")
CAMPO_DATA_CORRECAO = "dtCorrecao"
CAMPO_DATA_POSICAO = "dtPosicao"
IDENTIFICADORES_DATA_CORRECAO = (
    "dataCorrecao",
    "dtCorrecao",
    "filter.dtCorrecao",
    "filter.dtCorrecaoStr",
)
IDENTIFICADORES_DATA_POSICAO = (
    "dataPosicao",
    "dtPosicao",
    "filter.dtPosicao",
    "filter.dtPosicaoStr",
)
XPATH_EMPRESA_PRINCIPAL = "/html/body/div[1]/div/div[4]/main/div[1]/div[2]/div[1]/div/div[2]/div/div/div/div/div/form/div/div/div[2]/div/div/div/input"
XPATH_EMPRESA_ID_DINAMICO = "//*[@id='mui-58']"
XPATH_DATA_INICIO_PRINCIPAL = "/html/body/div[1]/div/div[4]/main/div[1]/div[2]/div[1]/div/div[2]/div/div/div/div/div/form/div/div/div[9]/div[1]/div/div/input"
XPATH_DATA_FIM_PRINCIPAL = "/html/body/div[1]/div/div[4]/main/div[1]/div[2]/div[1]/div/div[2]/div/div/div/div/div/form/div/div/div[9]/div[2]/div/div/input"
XPATH_ABRIR_FILTROS_DATAS = "//*[@id='main']/div[1]/div[2]/div[1]/div/div[2]/div/div/div/div/div/form/div/div/div[9]/div[5]/span/button"
XPATH_DATA_CORRECAO_MODAL = "/html/body/div[3]/div[3]/div/div[2]/div[2]/div[1]/div[1]/div/div/input"
XPATH_DATA_POSICAO_MODAL = "/html/body/div[3]/div[3]/div/div[2]/div[2]/div[1]/div[2]/div/div/input"
XPATH_APLICAR_FILTROS_MODAL = "/html/body/div[3]/div[3]/div/div[3]/button"
XPATH_CONSULTAR_PRINCIPAL = "/html/body/div[1]/div/div[4]/main/div[1]/div[2]/div[1]/div/div[2]/div/div/div/div/div/form/div/div/div[9]/div[7]/button"
XPATH_PAGINACAO_PRINCIPAL = "/html/body/div[1]/div/div[4]/main/div[1]/div[2]/div[2]/div/div[3]/div[2]/div/div[2]/div"
XPATH_PAGINACAO_5000_PRINCIPAL = "/html/body/div[3]/div[3]/ul/li[4]"
XPATH_GERAR_RELATORIO_PRINCIPAL = "//*[@id='main']/div[1]/div[2]/div[2]/div/div[1]/button[5]"
XPATH_FORMATO_EXPORTACAO_PRINCIPAL = "/html/body/div[4]/div[3]/div/div[2]/div/div[4]/div"
XPATH_OPCAO_EXCEL_PRINCIPAL = "//*[@id='mui-92']/li[2]"
XPATH_EXPORTAR_PRINCIPAL = "/html/body/div[4]/div[3]/div/div[3]/button[2]"
XPATH_PROXIMA_PAGINA_PRINCIPAL = "//*[@id='main']/div[1]/div[2]/div[2]/div/div[3]/div[2]/div/div[3]/button[2]"
XPATH_PROXIMA_PAGINA_ICONE = "//*[@id='main']/div[1]/div[2]/div[2]/div/div[3]/div[2]/div/div[3]/button[2]/svg"


class ModalDatasError(RuntimeError):
    etapa = "preenchimento_modal"


def _log_etapa(logger, iteracao, etapa, mensagem):
    print(f"[saldo_devedor][iteracao={iteracao}][etapa={etapa}] {mensagem}", flush=True)
    logger.info("[saldo_devedor][iteracao=%s][etapa=%s] %s", iteracao, etapa, mensagem)


def _wait_lento(driver):
    return WebDriverWait(driver, TIMEOUT_CARREGAMENTO_LENTO, poll_frequency=0.25)


def _aguardar_tela_filtros(driver, logger):
    driver.switch_to.default_content()
    espera = _wait_lento(driver)
    espera.until(EC.presence_of_element_located((By.XPATH, XPATH_DATA_INICIO_PRINCIPAL)))
    espera.until(EC.element_to_be_clickable((By.CSS_SELECTOR, "button#saldo-devedor[type='submit']")))
    logger.info("Formulario do saldo devedor disponivel para preenchimento.")
    print(
        "[saldo_devedor][diagnostico][espera] formulario pronto para preenchimento",
        flush=True,
    )


def _obter_texto_paginacao_atual(driver):
    driver.switch_to.default_content()
    elementos = driver.find_elements(
        By.XPATH,
        "//p[contains(@class,'MuiTablePagination-displayedRows')]",
    )
    if not elementos:
        return ""
    return (elementos[0].text or "").strip()


def _ha_indicador_carregamento_visivel(driver):
    localizadores = [
        (By.CSS_SELECTOR, "[role='progressbar']"),
        (By.CSS_SELECTOR, ".MuiCircularProgress-root"),
        (By.CSS_SELECTOR, ".MuiLinearProgress-root"),
    ]
    for localizador in localizadores:
        for elemento in driver.find_elements(*localizador):
            if elemento.is_displayed():
                return True
    return False


def _aguardar_grade_pronta(driver, logger, texto_anterior="", exigir_mudanca=False):
    driver.switch_to.default_content()
    time.sleep(PAUSA_ESTABILIZACAO)
    espera = _wait_lento(driver)
    if exigir_mudanca and texto_anterior:
        texto_atual = espera.until(
            lambda d: (
                texto
                if (texto := _obter_texto_paginacao_atual(d)) and texto != texto_anterior
                else False
            )
        )
    else:
        texto_atual = espera.until(lambda d: _obter_texto_paginacao_atual(d) or False)

    espera.until(lambda d: not _ha_indicador_carregamento_visivel(d))
    logger.info("Grade pronta para interacao | paginacao=%s", texto_atual)
    print(
        f"[saldo_devedor][diagnostico][espera] grade pronta | paginacao='{texto_atual}'",
        flush=True,
    )
    return texto_atual


def executar(driver, wait, relatorio_config, contexto):
    logger = contexto["logger"]
    normalizar_valor = contexto["normalizar_valor"]
    preencher_autocomplete_enter = contexto["preencher_autocomplete_enter"]
    marcar_inicio_download = contexto["marcar_inicio_download"]
    renomear_download = contexto["renomear_download"]
    sanitizar_nome_arquivo = contexto["sanitizar_nome_arquivo"]
    copiar_download_para_sharepoint = contexto["copiar_download_para_sharepoint"]
    montar_caminho_download_unico = contexto["montar_caminho_download_unico"]
    capture_debug_snapshot = contexto["capture_debug_snapshot"]
    download_dir = Path(contexto["download_dir"])

    filtros = relatorio_config.get("filtros", []) or []
    _log_filtros_recebidos(logger, filtros)
    empresa = _obter_filtro_por_campo(filtros, CAMPO_EMPRESA, *IDENTIFICADORES_EMPRESA)
    data_inicio = _obter_filtro_por_campo(filtros, CAMPO_DATA_INICIO, *IDENTIFICADORES_DATA_INICIO)
    data_fim = _obter_filtro_por_campo(filtros, CAMPO_DATA_FIM, *IDENTIFICADORES_DATA_FIM)
    data_correcao = _obter_filtro_por_campo(
        filtros,
        CAMPO_DATA_CORRECAO,
        *IDENTIFICADORES_DATA_CORRECAO,
    ) or _obter_filtro_por_nome(filtros, "data de correção") or _obter_filtro_por_nome(
        filtros, "data de correcao"
    ) or _obter_filtro_por_nome(filtros, "corre")
    data_posicao = _obter_filtro_por_campo(
        filtros,
        CAMPO_DATA_POSICAO,
        *IDENTIFICADORES_DATA_POSICAO,
    ) or _obter_filtro_por_nome(filtros, "data de posição") or _obter_filtro_por_nome(
        filtros, "data de posicao"
    ) or _obter_filtro_por_nome(filtros, "posi")

    if not empresa:
        empresa = _obter_filtro_por_nome(filtros, "empresa")
        if empresa:
            print(
                "[saldo_devedor][diagnostico] filtro de empresa localizado por nome como fallback.",
                flush=True,
            )
            logger.info("Filtro de empresa localizado por nome como fallback.")

    empresas = _extrair_valores(empresa)
    total = max(len(empresas), 1)
    valor_data_inicio = normalizar_valor(data_inicio) if data_inicio else ""
    valor_data_fim = normalizar_valor(data_fim) if data_fim else ""
    valor_data_correcao = normalizar_valor(data_correcao) if data_correcao else ""
    valor_data_posicao = normalizar_valor(data_posicao) if data_posicao else ""

    logger.info(
        (
            "Executor extrato_saldo_devedor_presente iniciando "
            "com empresa=%s | data_inicio=%s | data_fim=%s | "
            "data_correcao=%s | data_posicao=%s | iteracoes=%s."
        ),
        empresas or [""],
        valor_data_inicio,
        valor_data_fim,
        valor_data_correcao,
        valor_data_posicao,
        total,
    )

    for index in range(total):
        iteracao = index + 1
        valor_empresa = _obter_valor_iteracao(empresas, index)
        _log_etapa(logger, iteracao, "inicio", "iniciando fluxo do relatorio com filtros do formulario")
        _aguardar_tela_filtros(driver, logger)
        _log_etapa(logger, iteracao, "espera_inicial", "campos do formulario disponiveis; capturando snapshot")
        capture_debug_snapshot(f"saldo_devedor_inicio_iteracao_{iteracao}")

        _log_etapa(
            logger,
            iteracao,
            "preparacao",
            f"preenchendo filtros: empresa={valor_empresa}, data_inicio={valor_data_inicio}, data_fim={valor_data_fim}",
        )
        _preparar_filtros_relatorio(
            driver,
            wait,
            normalizar_valor,
            preencher_autocomplete_enter,
            empresa,
            valor_empresa,
            data_inicio,
            data_fim,
            data_correcao,
            data_posicao,
            logger,
            index,
            total,
            capture_debug_snapshot,
        )

        _log_etapa(logger, iteracao, "paginacao", f"ajustando paginacao para {TAMANHO_PAGINA}")
        _selecionar_tamanho_pagina(driver, wait, logger, TAMANHO_PAGINA)
        capture_debug_snapshot(f"saldo_devedor_paginacao_5000_iteracao_{iteracao}")
        total_registros = _obter_total_registros(driver, wait, logger)
        total_paginas = max(1, math.ceil(total_registros / int(TAMANHO_PAGINA)))
        _log_etapa(
            logger,
            iteracao,
            "paginacao",
            f"total de registros={total_registros}; total de paginas para exportacao={total_paginas}",
        )

        logger.info(
            "Saldo devedor presente retornou %s registro(s); executando %s exportacao(oes).",
            total_registros,
            total_paginas,
        )

        arquivos_baixados = []
        for pagina in range(total_paginas):
            numero_pagina = pagina + 1
            _log_etapa(
                logger,
                iteracao,
                "exportacao",
                f"iniciando exportacao da pagina {numero_pagina}/{total_paginas}",
            )
            logger.info(
                "Exportando pagina %s/%s do saldo devedor presente.",
                numero_pagina,
                total_paginas,
            )
            inicio_download = marcar_inicio_download()
            _clicar_botao_gerar_relatorio(driver, wait, logger)
            capture_debug_snapshot(f"saldo_devedor_modal_exportacao_iteracao_{iteracao}_pagina_{numero_pagina}")
            _selecionar_formato_excel_na_modal(driver, logger)
            time.sleep(PAUSA_INICIO_DOWNLOAD)

            arquivo_baixado = renomear_download(
                relatorio_config,
                _montar_sufixo_exportacao(valor_empresa, numero_pagina),
                inicio_download,
                registrar_resultado=False,
            )
            if arquivo_baixado:
                arquivos_baixados.append(Path(arquivo_baixado))
                _log_etapa(
                    logger,
                    iteracao,
                    "download",
                    f"arquivo exportado com sucesso: {arquivo_baixado}",
                )
            else:
                _log_etapa(
                    logger,
                    iteracao,
                    "download",
                    f"nenhum arquivo encontrado para a pagina {numero_pagina}",
                )

            if pagina < total_paginas - 1:
                _log_etapa(logger, iteracao, "paginacao", f"avancando para a proxima pagina apos exportar a pagina {numero_pagina}")
                _ir_para_proxima_pagina(driver, logger)

        _log_etapa(logger, iteracao, "consolidacao", f"{len(arquivos_baixados)} arquivo(s) exportado(s)")
        _registrar_arquivo_final(
            arquivos_baixados,
            relatorio_config,
            valor_empresa,
            download_dir,
            sanitizar_nome_arquivo,
            montar_caminho_download_unico,
            copiar_download_para_sharepoint,
            logger,
        )


def _registrar_arquivo_final(
    arquivos,
    relatorio_config,
    empresa_id,
    download_dir,
    sanitizar_nome_arquivo,
    montar_caminho_download_unico,
    copiar_download_para_sharepoint,
    logger,
):
    """Copia para o destino e registra so o resultado da iteracao.

    Uma pagina: o proprio arquivo. Varias: o consolidado, e as paginas locais sao
    apagadas. Se a consolidacao nao gerar linhas, as paginas exportadas ficam como resultado.
    """
    if len(arquivos) == 1:
        copiar_download_para_sharepoint(relatorio_config, arquivos[0])
        return
    if not arquivos:
        return

    consolidado = _consolidar_planilhas(
        arquivos,
        relatorio_config,
        empresa_id,
        download_dir,
        sanitizar_nome_arquivo,
        montar_caminho_download_unico,
        logger,
    )
    if consolidado:
        logger.info("Arquivo consolidado gerado: %s", consolidado)
        copiar_download_para_sharepoint(relatorio_config, consolidado)
        _remover_arquivos_intermediarios(arquivos, consolidado, logger)
        return

    for arquivo in arquivos:
        copiar_download_para_sharepoint(relatorio_config, arquivo)


def _preparar_filtros_relatorio(
    driver,
    wait,
    normalizar_valor,
    preencher_autocomplete_enter,
    filtro_empresa,
    valor_empresa,
    filtro_data_inicio,
    filtro_data_fim,
    filtro_data_correcao,
    filtro_data_posicao,
    logger,
    index,
    total,
    capture_debug_snapshot,
):
    iteracao = index + 1
    if filtro_data_inicio:
        valor_data_inicio = normalizar_valor(filtro_data_inicio)
        _log_etapa(logger, iteracao, "data_inicio", f"preenchendo data inicial com {valor_data_inicio}")
        valor_data_inicio_final = _preencher_input_manual(
            driver,
            wait,
            _montar_identificadores(
                filtro_data_inicio,
                XPATH_DATA_INICIO_PRINCIPAL,
                *IDENTIFICADORES_DATA_INICIO,
            ),
            valor_data_inicio,
            logger,
        )
        _log_etapa(
            logger,
            iteracao,
            "data_inicio",
            f"campo data inicial retornou '{valor_data_inicio_final}'",
        )
        capture_debug_snapshot(f"saldo_devedor_data_inicio_iteracao_{iteracao}")

    if filtro_data_fim:
        valor_data_fim = normalizar_valor(filtro_data_fim)
        _log_etapa(logger, iteracao, "data_fim", f"preenchendo data final com {valor_data_fim}")
        valor_data_fim_final = _preencher_input_manual(
            driver,
            wait,
            _montar_identificadores(
                filtro_data_fim,
                XPATH_DATA_FIM_PRINCIPAL,
                *IDENTIFICADORES_DATA_FIM,
            ),
            valor_data_fim,
            logger,
        )
        _log_etapa(
            logger,
            iteracao,
            "data_fim",
            f"campo data final retornou '{valor_data_fim_final}'",
        )
        capture_debug_snapshot(f"saldo_devedor_data_fim_iteracao_{iteracao}")

    if filtro_empresa and valor_empresa:
        _log_etapa(logger, iteracao, "empresa", f"preenchendo empresa autocomplete com {valor_empresa}")
        valor_final = _preencher_autocomplete_empresa(
            driver,
            wait,
            valor_empresa,
            logger,
            _montar_identificadores(
                filtro_empresa,
                XPATH_EMPRESA_ID_DINAMICO,
                XPATH_EMPRESA_PRINCIPAL,
                *IDENTIFICADORES_EMPRESA,
            ),
        )
        if not valor_final:
            try:
                valor_final = preencher_autocomplete_enter(
                    driver,
                    wait,
                    _obter_identificador_principal(filtro_empresa),
                    valor_empresa,
                    placeholder="Pesquisar empresa",
                    abrir_dropdown=True,
                )
            except Exception as exc:
                print(
                    f"[saldo_devedor][diagnostico][empresa] fallback generico falhou: {type(exc).__name__}",
                    flush=True,
                )
                logger.info(
                    "Fallback generico da empresa nao conseguiu preencher o campo: %s.",
                    type(exc).__name__,
                )
        logger.info(
            "Empresa autocomplete preenchida com ENTER | valor=%s | valor_final=%s",
            valor_empresa,
            valor_final,
        )
        _log_etapa(
            logger,
            iteracao,
            "empresa",
            f"campo empresa retornou '{valor_final}'",
        )
        capture_debug_snapshot(f"saldo_devedor_empresa_preenchida_iteracao_{iteracao}")

    primeira_iteracao_com_datas_modal = index == 0 and any(
        normalizar_valor(filtro)
        for filtro in (filtro_data_correcao, filtro_data_posicao)
        if filtro
    )
    if primeira_iteracao_com_datas_modal:
        _log_etapa(
            logger,
            iteracao,
            "datas_modal",
            "abrindo modal para preencher as datas de correcao e de posicao por ultimo",
        )
        _aplicar_datas_modal(
            driver,
            wait,
            normalizar_valor,
            filtro_data_correcao,
            filtro_data_posicao,
            logger,
        )
        capture_debug_snapshot(f"saldo_devedor_datas_modal_aplicadas_iteracao_{iteracao}")
        _log_etapa(
            logger,
            iteracao,
            "datas_modal",
            "filtros aplicados pela modal; primeira consulta disparada automaticamente",
        )
    else:
        _log_etapa(logger, iteracao, "consultar", "clicando em Consultar com vencimento e empresa preenchidos")
        logger.info(
            "Executor extrato_saldo_devedor_presente clicando em Consultar com vencimento e empresa preenchidos na iteracao %s/%s.",
            iteracao,
            total,
        )
        _clicar_botao_consultar(driver, wait, logger)

    _aguardar_grade_pronta(driver, logger)
    capture_debug_snapshot(f"saldo_devedor_grade_carregada_iteracao_{iteracao}")
    _log_etapa(logger, iteracao, "consultar", "grade carregada apos executar a pesquisa")


def _aplicar_datas_modal(
    driver,
    wait,
    normalizar_valor,
    filtro_data_correcao,
    filtro_data_posicao,
    logger,
):
    driver.switch_to.default_content()
    botao_abrir = wait.until(
        EC.element_to_be_clickable((By.XPATH, XPATH_ABRIR_FILTROS_DATAS))
    )
    driver.execute_script("arguments[0].scrollIntoView({block: 'center'});", botao_abrir)
    try:
        botao_abrir.click()
    except Exception:
        driver.execute_script("arguments[0].click();", botao_abrir)

    # Os Dialogs do Material UI sao montados em um portal. O indice do div no
    # body muda conforme menus, tooltips e outros portais ja abertos, portanto
    # os XPaths absolutos antigos deixavam de encontrar campos que estavam
    # visiveis. Aguarde a modal pelo seu papel semantico.
    try:
        wait.until(
            lambda d: any(
                elemento.is_displayed()
                for elemento in d.find_elements(
                    By.CSS_SELECTOR,
                    "[role='dialog'], [aria-modal='true'], .MuiDialog-root",
                )
            )
        )
    except TimeoutException as exc:
        raise ModalDatasError("Modal de datas: a janela de filtros nao ficou disponivel.") from exc

    if filtro_data_correcao:
        valor = normalizar_valor(filtro_data_correcao)
        _preencher_input_modal(
            driver,
            wait,
            XPATH_DATA_CORRECAO_MODAL,
            valor,
            logger,
            "data_correcao",
            rotulo="Correcao ate",
            indice_fallback=0,
        )

    if filtro_data_posicao:
        valor = normalizar_valor(filtro_data_posicao)
        _preencher_input_modal(
            driver,
            wait,
            XPATH_DATA_POSICAO_MODAL,
            valor,
            logger,
            "data_posicao",
            rotulo="Posicao em",
            indice_fallback=1,
        )

    try:
        botao_aplicar = wait.until(_localizar_botao_aplicar_modal)
    except TimeoutException as exc:
        raise ModalDatasError("Modal de datas: o botao 'Aplicar filtros' nao foi localizado.") from exc
    try:
        botao_aplicar.click()
    except Exception:
        driver.execute_script("arguments[0].click();", botao_aplicar)
    logger.info("Botao 'Aplicar filtros' da modal clicado; consulta inicial disparada.")


def _elemento_interagivel(elemento):
    try:
        return elemento.is_displayed() and elemento.is_enabled()
    except Exception:
        return False


def _inputs_visiveis_modal(driver):
    seletores = (
        "[role='dialog'] input",
        "[aria-modal='true'] input",
        ".MuiDialog-root input",
    )
    encontrados = []
    ids = set()
    for seletor in seletores:
        for elemento in driver.find_elements(By.CSS_SELECTOR, seletor):
            identidade = getattr(elemento, "id", None) or id(elemento)
            if identidade not in ids and _elemento_interagivel(elemento):
                ids.add(identidade)
                encontrados.append(elemento)
    return encontrados


def _localizar_input_modal(driver, nome_campo, xpath_legado, rotulo, indice_fallback):
    termos = {
        "data_correcao": ("Correcao ate", "Correção até", "correcao", "correção"),
        "data_posicao": ("Posicao em", "Posição em", "posicao", "posição"),
    }.get(nome_campo, (rotulo,))

    raizes = ("//*[@role='dialog']", "//*[@aria-modal='true']", "//*[contains(@class,'MuiDialog-root')]")
    for raiz in raizes:
        for termo in termos:
            localizadores = (
                (
                    By.XPATH,
                    f"{raiz}//*[self::label or self::legend][contains(normalize-space(.), '{termo}')]/ancestor::div[contains(@class,'MuiFormControl-root')][1]//input",
                ),
                (
                    By.XPATH,
                    f"{raiz}//input[contains(@aria-label, '{termo}') or contains(@placeholder, '{termo}') or contains(@name, '{termo}') or contains(@id, '{termo}')]",
                ),
            )
            for localizador in localizadores:
                for elemento in driver.find_elements(*localizador):
                    if _elemento_interagivel(elemento):
                        return elemento

    inputs = _inputs_visiveis_modal(driver)
    if len(inputs) > indice_fallback:
        return inputs[indice_fallback]

    # Compatibilidade com uma eventual tela antiga que ainda use a estrutura
    # fixa. Esse e apenas o ultimo recurso, nunca o seletor principal.
    for elemento in driver.find_elements(By.XPATH, xpath_legado):
        if _elemento_interagivel(elemento):
            return elemento
    return False


def _localizar_botao_aplicar_modal(driver):
    texto_minusculo = "translate(normalize-space(.), 'ABCDEFGHIJKLMNOPQRSTUVWXYZ', 'abcdefghijklmnopqrstuvwxyz')"
    localizadores = (
        (By.XPATH, f"//*[@role='dialog']//button[contains({texto_minusculo}, 'aplicar')]"),
        (By.XPATH, f"//*[@aria-modal='true']//button[contains({texto_minusculo}, 'aplicar')]"),
        (By.XPATH, f"//*[contains(@class,'MuiDialog-root')]//button[contains({texto_minusculo}, 'aplicar')]"),
        (By.XPATH, XPATH_APLICAR_FILTROS_MODAL),
    )
    for localizador in localizadores:
        for elemento in driver.find_elements(*localizador):
            if _elemento_interagivel(elemento):
                return elemento
    return False


def _apenas_digitos(valor):
    return re.sub(r"\D", "", str(valor or ""))


def _preencher_input_modal(
    driver,
    wait,
    xpath,
    valor,
    logger,
    nome_campo,
    rotulo="",
    indice_fallback=0,
):
    try:
        campo = wait.until(
            lambda d: _localizar_input_modal(
                d,
                nome_campo,
                xpath,
                rotulo,
                indice_fallback,
            )
        )
    except TimeoutException as exc:
        nome_legivel = rotulo or nome_campo
        raise ModalDatasError(f"Modal de datas: o campo '{nome_legivel}' nao foi localizado.") from exc
    driver.execute_script("arguments[0].scrollIntoView({block: 'center'});", campo)
    try:
        campo.click()
        campo.clear()
        campo.send_keys(Keys.CONTROL, "a")
        campo.send_keys(Keys.DELETE)
        campo.send_keys(valor)
        campo.send_keys(Keys.TAB)
    except Exception:
        _definir_valor_controlado(driver, campo, valor)

    valor_final = (campo.get_attribute("value") or campo.text or "").strip()
    if _apenas_digitos(valor_final) != _apenas_digitos(valor):
        _definir_valor_controlado(driver, campo, valor)
        valor_final = (campo.get_attribute("value") or campo.text or "").strip()
    if _apenas_digitos(valor_final) != _apenas_digitos(valor):
        nome_legivel = rotulo or nome_campo
        raise ModalDatasError(
            f"Modal de datas: o campo '{nome_legivel}' nao aceitou a data informada."
        )
    logger.info(
        "Campo %s da modal preenchido | valor=%s | valor_final=%s",
        nome_campo,
        valor,
        valor_final,
    )
    return valor_final


def _preencher_input_manual(driver, wait, identificadores, valor, logger):
    identificadores = tuple(
        item for item in (identificadores if isinstance(identificadores, (list, tuple)) else [identificadores]) if item
    )
    campo, iframe_index = _localizar_campo_em_iframes(driver, wait, identificadores)
    try:
        identificador_resolvido = campo.get_attribute("id") or campo.get_attribute("name") or ""
        valor_antes = (campo.get_attribute("value") or campo.text or "").strip()
        print(
            f"[saldo_devedor][diagnostico][campo_manual] iframe={iframe_index} | identificadores={list(identificadores)} | resolvido={identificador_resolvido} | valor_antes='{valor_antes}'",
            flush=True,
        )
        logger.info(
            "Campo manual localizado | iframe=%s | identificadores=%s | resolvido=%s | valor_antes=%s",
            iframe_index,
            list(identificadores),
            identificador_resolvido,
            valor_antes,
        )

        driver.execute_script("arguments[0].scrollIntoView({block: 'center'});", campo)
        time.sleep(0.2)

        try:
            campo.click()
            time.sleep(0.3)
        except Exception:
            pass

        try:
            campo.clear()
            time.sleep(0.3)
            campo.send_keys(Keys.CONTROL, "a")
            time.sleep(0.1)
            campo.send_keys(Keys.DELETE)
            time.sleep(0.2)
            campo.send_keys(valor)
            time.sleep(0.3)
            campo.send_keys(Keys.TAB)
            print(
                f"[saldo_devedor][diagnostico][campo_manual] send_keys executado para {identificador_resolvido} com valor '{valor}'",
                flush=True,
            )
        except Exception:
            _definir_valor_controlado(driver, campo, valor)
            print(
                f"[saldo_devedor][diagnostico][campo_manual] fallback JS executado para {identificador_resolvido} com valor '{valor}'",
                flush=True,
            )
        time.sleep(0.8)
        valor_final = (campo.get_attribute("value") or campo.text or "").strip()
        logger.info(
            "Campo manual preenchido | iframe=%s | campo=%s | valor=%s | valor_final=%s",
            iframe_index,
            identificador_resolvido,
            valor,
            valor_final,
        )
        print(
            f"[saldo_devedor][diagnostico][campo_manual] iframe={iframe_index} | resolvido={identificador_resolvido} | valor_final='{valor_final}'",
            flush=True,
        )
        return valor_final
    finally:
        driver.switch_to.default_content()


def _fechar_janelas_extras(driver):
    main_window_handle = driver.current_window_handle
    handles = list(driver.window_handles)
    for handle in handles:
        if handle == main_window_handle:
            continue
        driver.switch_to.window(handle)
        driver.close()
    driver.switch_to.window(main_window_handle)


def _com_iframe(driver, wait):
    driver.switch_to.default_content()
    iframes = wait.until(EC.presence_of_all_elements_located((By.TAG_NAME, "iframe")))
    print(
        f"[saldo_devedor][diagnostico][iframe] total_iframes={len(iframes)} | usando_indice=0",
        flush=True,
    )
    iframe = iframes[0]
    driver.switch_to.frame(iframe)


def _localizar_campo_em_iframes(driver, wait, identificadores):
    driver.switch_to.default_content()
    iframes = wait.until(EC.presence_of_all_elements_located((By.TAG_NAME, "iframe")))
    print(
        f"[saldo_devedor][diagnostico][iframe] total_iframes={len(iframes)} | buscando_campo_em_todos",
        flush=True,
    )

    ultimo_erro = None
    contextos = [("default", None)] + [(indice, iframe) for indice, iframe in enumerate(iframes)]

    for contexto, iframe in contextos:
        driver.switch_to.default_content()
        try:
            if iframe is not None:
                driver.switch_to.frame(iframe)

            campo = _localizar_campo_no_contexto_atual(driver, identificadores)
            if campo is not None:
                return campo, contexto
        except Exception as exc:
            ultimo_erro = exc

    driver.switch_to.default_content()
    identificadores_texto = ", ".join(identificadores)
    raise ultimo_erro or NoSuchElementException(
        f"Campo nao localizado em nenhum iframe. Tentativas: {identificadores_texto}"
    )


def _preencher_autocomplete_empresa(driver, wait, valor, logger, identificadores):
    driver.switch_to.default_content()
    iframes = wait.until(EC.presence_of_all_elements_located((By.TAG_NAME, "iframe")))
    ultimo_erro = None

    contextos = [("default", None)] + [(indice, iframe) for indice, iframe in enumerate(iframes)]

    for contexto, iframe in contextos:
        driver.switch_to.default_content()
        try:
            if iframe is not None:
                driver.switch_to.frame(iframe)
            campo = _localizar_autocomplete_empresa_no_iframe(driver, wait, identificadores)
            driver.execute_script("arguments[0].scrollIntoView({block: 'center'});", campo)
            time.sleep(0.2)
            valor_final = _executar_fluxo_autocomplete_empresa(driver, wait, campo, valor)
            time.sleep(0.5)
            valor_final = (campo.get_attribute("value") or "").strip()
            print(
                f"[saldo_devedor][diagnostico][empresa] iframe={contexto} | valor_enviado='{valor}' | valor_final='{valor_final}'",
                flush=True,
            )
            logger.info(
                "Autocomplete empresa preenchido | iframe=%s | valor_enviado=%s | valor_final=%s",
                contexto,
                valor,
                valor_final,
            )
            if valor_final:
                return valor_final
        except Exception as exc:
            ultimo_erro = exc
        finally:
            driver.switch_to.default_content()

    if ultimo_erro:
        logger.warning("Falha ao preencher autocomplete da empresa em todos os iframes.")
    return ""


def _localizar_autocomplete_empresa_no_iframe(driver, wait, identificadores):
    localizadores = [
        (By.XPATH, XPATH_EMPRESA_ID_DINAMICO),
        (By.XPATH, XPATH_EMPRESA_PRINCIPAL),
        (By.XPATH, "//input[@role='combobox' and @placeholder='Pesquisar empresa']"),
        (By.XPATH, "//input[@role='combobox' and contains(@class,'MuiAutocomplete-input')]"),
    ]

    for identificador in identificadores:
        localizadores_identificador = _montar_localizadores_identificador(identificador)
        for localizador in reversed(localizadores_identificador):
            localizadores.insert(0, localizador)

    ultimo_erro = None
    for localizador in localizadores:
        try:
            return WebDriverWait(driver, TIMEOUT_LOCALIZACAO_CAMPO).until(
                EC.element_to_be_clickable(localizador)
            )
        except Exception as exc:
            ultimo_erro = exc

    raise ultimo_erro or TimeoutException("Nao foi possivel localizar o autocomplete da empresa.")


def _executar_fluxo_autocomplete_empresa(driver, wait, campo, valor):
    valor = str(valor or "").strip()
    if not valor:
        return ""

    try:
        campo.click()
        time.sleep(0.2)
    except Exception:
        driver.execute_script("arguments[0].click();", campo)
        time.sleep(0.2)

    try:
        campo.send_keys(Keys.CONTROL, "a")
        time.sleep(0.1)
        campo.send_keys(Keys.DELETE)
        time.sleep(0.2)
        campo.send_keys(valor)
    except Exception:
        _definir_valor_controlado(driver, campo, valor)

    time.sleep(1)

    if _selecionar_opcao_autocomplete_empresa(driver, wait, valor):
        return (campo.get_attribute("value") or "").strip()

    try:
        campo.send_keys(Keys.ARROW_DOWN)
        time.sleep(0.3)
        campo.send_keys(Keys.ENTER)
        time.sleep(1)
    except Exception:
        pass

    valor_final = (campo.get_attribute("value") or "").strip()
    if valor_final:
        return valor_final

    _definir_valor_controlado(driver, campo, valor)
    time.sleep(0.5)
    return (campo.get_attribute("value") or "").strip()


def _selecionar_opcao_autocomplete_empresa(driver, wait, valor):
    termo = str(valor or "").strip()
    if not termo:
        return False

    localizadores = [
        (By.XPATH, f"//li[@role='option' and contains(normalize-space(.), '{termo}')]"),
        (By.XPATH, f"//div[@role='option' and contains(normalize-space(.), '{termo}')]"),
        (By.XPATH, "//li[@role='option']"),
        (By.XPATH, "//div[@role='option']"),
    ]

    for localizador in localizadores:
        try:
            opcoes = WebDriverWait(driver, TIMEOUT_LOCALIZACAO_CAMPO).until(
                EC.presence_of_all_elements_located(localizador)
            )
            for opcao in opcoes:
                try:
                    if not opcao.is_displayed():
                        continue
                    driver.execute_script(
                        "arguments[0].scrollIntoView({block: 'nearest'});",
                        opcao,
                    )
                    time.sleep(0.1)
                    opcao.click()
                    time.sleep(1)
                    return True
                except Exception:
                    continue
        except Exception:
            continue

    return False


def _definir_valor_controlado(driver, campo, valor):
    driver.execute_script(
        """
        const campo = arguments[0];
        const valor = arguments[1];
        const proto = Object.getPrototypeOf(campo);
        const descriptor = Object.getOwnPropertyDescriptor(proto, 'value')
            || Object.getOwnPropertyDescriptor(HTMLInputElement.prototype, 'value')
            || Object.getOwnPropertyDescriptor(HTMLTextAreaElement.prototype, 'value');
        campo.focus();
        if (descriptor && descriptor.set) {
            descriptor.set.call(campo, valor);
        } else if ('value' in campo) {
            campo.value = valor;
        } else {
            campo.textContent = valor;
        }
        campo.dispatchEvent(new Event('input', { bubbles: true }));
        campo.dispatchEvent(new Event('change', { bubbles: true }));
        campo.dispatchEvent(new KeyboardEvent('keydown', { bubbles: true, key: 'Tab' }));
        campo.dispatchEvent(new KeyboardEvent('keyup', { bubbles: true, key: 'Tab' }));
        campo.blur();
        """,
        campo,
        valor,
    )


def _localizar_campo_no_contexto_atual(driver, identificadores):
    for identificador in identificadores:
        for localizador in _montar_localizadores_identificador(identificador):
            elementos = driver.find_elements(*localizador)
            if elementos:
                return elementos[0]
    return None


def _montar_localizadores_identificador(identificador):
    identificador = str(identificador or "").strip()
    if not identificador:
        return []

    if (
        identificador.startswith(("/", "("))
        or identificador.startswith("./")
        or identificador.startswith("//")
        or identificador.startswith("//*[@")
    ):
        return [(By.XPATH, identificador)]

    return [
        (By.ID, identificador),
        (By.NAME, identificador),
        (By.XPATH, f"//*[@id='{identificador}' or @name='{identificador}']"),
    ]


def _clicar_botao_consultar(driver, wait, logger):
    driver.switch_to.default_content()
    iframes = driver.find_elements(By.TAG_NAME, "iframe")
    contextos = [("default", None)] + [(indice, iframe) for indice, iframe in enumerate(iframes)]
    localizadores = [
        (By.CSS_SELECTOR, "button#saldo-devedor[type='submit']"),
        (By.ID, "saldo-devedor"),
        (By.XPATH, "//*[@id='saldo-devedor']"),
        (By.XPATH, f"{XPATH_CONSULTAR_PRINCIPAL}[normalize-space(.)='Consultar']"),
        (By.XPATH, XPATH_CONSULTAR_PRINCIPAL),
    ]
    ultimo_erro = None

    try:
        for contexto, iframe in contextos:
            driver.switch_to.default_content()
            if iframe is not None:
                driver.switch_to.frame(iframe)

            for localizador in localizadores:
                try:
                    botao = WebDriverWait(driver, TIMEOUT_LOCALIZACAO_CAMPO).until(
                        EC.element_to_be_clickable(localizador)
                    )
                    driver.execute_script("arguments[0].scrollIntoView({block: 'center'});", botao)
                    time.sleep(0.2)
                    try:
                        botao.click()
                    except Exception:
                        driver.execute_script("arguments[0].click();", botao)
                    time.sleep(0.5)
                    print(
                        f"[saldo_devedor][diagnostico][consultar] botao clicado | contexto={contexto} | locator={localizador}",
                        flush=True,
                    )
                    logger.info("Botao 'Consultar' clicado com sucesso no contexto %s.", contexto)
                    return
                except Exception as exc:
                    ultimo_erro = exc

        driver.switch_to.default_content()
        botao = WebDriverWait(driver, TIMEOUT_LOCALIZACAO_CAMPO).until(
            EC.presence_of_element_located((By.CSS_SELECTOR, "button#saldo-devedor[type='submit']"))
        )
        driver.execute_script(
            """
            const button = arguments[0];
            const form = button.closest('form');
            if (form && typeof form.requestSubmit === 'function') {
                form.requestSubmit(button);
            } else {
                button.dispatchEvent(new MouseEvent('click', { bubbles: true, cancelable: true }));
            }
            """,
            botao,
        )
        print(
            "[saldo_devedor][diagnostico][consultar] formulario submetido por fallback JS | contexto=default",
            flush=True,
        )
        logger.info("Formulario do saldo devedor submetido por fallback JS.")
    finally:
        driver.switch_to.default_content()


def _selecionar_tamanho_pagina(driver, wait, logger, valor_pagina):
    driver.switch_to.default_content()
    try:
        try:
            combobox = wait.until(
                EC.element_to_be_clickable(
                    (
                        By.XPATH,
                        "//div[@role='combobox' and contains(@class,'MuiTablePagination-select')]",
                    )
                )
            )
        except TimeoutException:
            combobox = wait.until(
                EC.element_to_be_clickable((By.XPATH, XPATH_PAGINACAO_PRINCIPAL))
            )
        valor_atual = (combobox.text or "").strip()
        if valor_atual != valor_pagina:
            combobox.click()
            try:
                opcao = wait.until(
                    EC.element_to_be_clickable(
                        (
                            By.XPATH,
                            f"//li[@role='option' and (@data-value='{valor_pagina}' or normalize-space(.)='{valor_pagina}')]",
                        )
                    )
                )
            except TimeoutException:
                opcao = wait.until(
                    EC.element_to_be_clickable((By.XPATH, XPATH_PAGINACAO_5000_PRINCIPAL))
                )
            opcao.click()
            print(
                f"[saldo_devedor][diagnostico][paginacao] valor selecionado='{valor_pagina}'",
                flush=True,
            )
            _aguardar_grade_pronta(driver, logger)
        else:
            print(
                f"[saldo_devedor][diagnostico][paginacao] valor ja selecionado='{valor_pagina}'",
                flush=True,
            )
        logger.info("Paginacao ajustada para %s itens por pagina.", valor_pagina)
    finally:
        driver.switch_to.default_content()


def _obter_total_registros(driver, wait, logger):
    driver.switch_to.default_content()
    try:
        texto = wait.until(
            EC.presence_of_element_located(
                (By.XPATH, "//p[contains(@class,'MuiTablePagination-displayedRows')]")
            )
        ).text
        correspondencia = re.search(r"de\s+([\d\.]+)", texto)
        if not correspondencia:
            raise ValueError(f"Nao foi possivel extrair o total de registros do texto '{texto}'.")
        total = int(correspondencia.group(1).replace(".", ""))
        logger.info("Total de registros identificado na grade: %s.", total)
        print(
            f"[saldo_devedor][diagnostico][paginacao] total_registros={total}",
            flush=True,
        )
        return total
    finally:
        driver.switch_to.default_content()


def _clicar_botao_gerar_relatorio(driver, wait, logger):
    driver.switch_to.default_content()
    try:
        try:
            botao = wait.until(
                EC.element_to_be_clickable((By.XPATH, XPATH_GERAR_RELATORIO_PRINCIPAL))
            )
        except TimeoutException:
            botao = wait.until(
                EC.element_to_be_clickable(
                    (
                        By.XPATH,
                        "//button[normalize-space(.)='Gerar Relatório' or normalize-space(.)='Gerar Relatorio' "
                        "or .//*[normalize-space(.)='Gerar Relatório'] or .//*[normalize-space(.)='Gerar Relatorio']]",
                    )
                )
            )
        botao.click()
        logger.info("Botao 'Gerar Relatorio' clicado com sucesso.")
        print(
            "[saldo_devedor][diagnostico][exportacao] botao 'Gerar Relatorio' clicado",
            flush=True,
        )
    finally:
        driver.switch_to.default_content()


def _selecionar_formato_excel_na_modal(driver, logger):
    driver.switch_to.default_content()
    espera_rapida = WebDriverWait(driver, TIMEOUT_LOCALIZACAO_CAMPO)
    espera = _wait_lento(driver)
    try:
        try:
            combobox = espera_rapida.until(
                EC.element_to_be_clickable((By.XPATH, XPATH_FORMATO_EXPORTACAO_PRINCIPAL))
            )
        except TimeoutException:
            combobox = espera.until(
                EC.element_to_be_clickable(
                    (
                        By.XPATH,
                        "//div[@role='combobox' and contains(@class,'MuiSelect-select') "
                        "and not(contains(@class,'MuiTablePagination-select'))]",
                    )
                )
            )
        combobox.click()
        try:
            opcao_excel = espera_rapida.until(
                EC.element_to_be_clickable((By.XPATH, XPATH_OPCAO_EXCEL_PRINCIPAL))
            )
        except TimeoutException:
            opcao_excel = espera.until(
                EC.element_to_be_clickable(
                    (By.XPATH, "//li[@role='option' and @data-value='excel']")
                )
            )
        opcao_excel.click()
        logger.info("Formato EXCEL selecionado na modal de exportacao.")
        print(
            "[saldo_devedor][diagnostico][exportacao] formato Excel selecionado na modal",
            flush=True,
        )
        try:
            botao_exportar = espera_rapida.until(
                EC.element_to_be_clickable((By.XPATH, XPATH_EXPORTAR_PRINCIPAL))
            )
        except TimeoutException:
            botao_exportar = espera.until(
                EC.element_to_be_clickable(
                    (By.XPATH, "//button[normalize-space(.)='Exportar']")
                )
            )
        botao_exportar.click()
        logger.info("Botao 'Exportar' clicado com sucesso na modal.")
        print(
            "[saldo_devedor][diagnostico][exportacao] botao 'Exportar' clicado na modal",
            flush=True,
        )
    except TimeoutException:
        logger.exception("Nao foi possivel concluir a modal de exportacao do saldo devedor.")
        print(
            "[saldo_devedor][diagnostico][exportacao] falha ao concluir a modal de exportacao",
            flush=True,
        )
        raise
    finally:
        driver.switch_to.default_content()


def _ir_para_proxima_pagina(driver, logger):
    driver.switch_to.default_content()
    texto_anterior = _obter_texto_paginacao_atual(driver)
    espera = WebDriverWait(driver, TIMEOUT_LOCALIZACAO_CAMPO)
    try:
        localizadores = [
            (
                By.XPATH,
                XPATH_PROXIMA_PAGINA_PRINCIPAL,
            ),
            (
                By.XPATH,
                XPATH_PROXIMA_PAGINA_ICONE,
            ),
            (
                By.XPATH,
                "//button[@aria-label='Go to next page' or @aria-label='Ir para a próxima página' "
                "or @aria-label='Ir para a proxima pagina' or @title='Go to next page']",
            ),
            (
                By.XPATH,
                "//button[.//*[@data-testid='KeyboardArrowRightIcon']]",
            ),
            (
                By.XPATH,
                "(//button[contains(@class,'MuiTablePagination-actions') or @type='button'])[last()]",
            ),
        ]

        ultimo_erro = None
        for localizador in localizadores:
            try:
                botao = espera.until(EC.element_to_be_clickable(localizador))
            except Exception as exc:
                ultimo_erro = exc
                continue

            if botao.tag_name.lower() == "svg":
                botao = botao.find_element(By.XPATH, "./ancestor::button[1]")
            driver.execute_script("arguments[0].scrollIntoView({block: 'center'});", botao)
            botao.click()
            logger.info("Avancando para a proxima pagina da grade.")
            print(
                f"[saldo_devedor][diagnostico][paginacao] proxima pagina acionada | locator={localizador}",
                flush=True,
            )
            # Depois do clique, uma falha na espera nao pode provocar outro clique:
            # isso saltaria a pagina seguinte e poderia exportar a pagina errada.
            _aguardar_grade_pronta(
                driver,
                logger,
                texto_anterior=texto_anterior,
                exigir_mudanca=True,
            )
            return

        if ultimo_erro:
            raise ultimo_erro
    finally:
        driver.switch_to.default_content()


def _montar_sufixo_exportacao(empresa_id, pagina):
    empresa = str(empresa_id or "sem_empresa").strip() or "sem_empresa"
    return f"{empresa}_parte_{pagina}"


def _obter_filtro_por_campo(filtros, *campos_identificadores):
    campos_normalizados = {
        _normalizar_texto(campo)
        for campo in campos_identificadores
        if str(campo or "").strip()
    }

    for filtro in filtros:
        campo = _normalizar_texto(filtro.get("campo_identificador"))
        if campo in campos_normalizados:
            return filtro

        if any(campo.endswith(f".{campo_esperado}") for campo_esperado in campos_normalizados):
            return filtro

    return None


def _obter_filtro_por_nome(filtros, trecho_nome):
    trecho = _normalizar_texto(trecho_nome)
    for filtro in filtros:
        nome = _normalizar_texto(filtro.get("nome") or filtro.get("nome_filtro"))
        if trecho and trecho in nome:
            return filtro
    return None


def _montar_identificadores(filtro, *fallbacks):
    identificadores = []
    campo_filtro = str(filtro.get("campo_identificador") or "").strip()
    if campo_filtro:
        identificadores.append(campo_filtro)

    for fallback in fallbacks:
        fallback = str(fallback or "").strip()
        if fallback and fallback not in identificadores:
            identificadores.append(fallback)

    return identificadores


def _obter_identificador_principal(filtro):
    identificadores = _montar_identificadores(filtro)
    return identificadores[0] if identificadores else ""


def _extrair_valores(filtro):
    if not filtro:
        return []

    valor = filtro.get("valor")
    if isinstance(valor, list):
        return [str(item).strip() for item in valor if str(item).strip()]

    return [parte.strip() for parte in str(valor or "").split(",") if parte.strip()]


def _obter_valor_iteracao(valores, index):
    if not valores:
        return ""

    if index < len(valores):
        return valores[index]

    if len(valores) == 1:
        return valores[0]

    return ""


def _normalizar_texto(valor):
    return str(valor or "").strip().lower()


def _log_filtros_recebidos(logger, filtros):
    if not filtros:
        print("[saldo_devedor][diagnostico] nenhum filtro recebido no relatorio.", flush=True)
        logger.info("Nenhum filtro recebido no relatorio de saldo devedor presente.")
        return

    print(
        f"[saldo_devedor][diagnostico] {len(filtros)} filtro(s) recebido(s) para o relatorio.",
        flush=True,
    )
    logger.info("%s filtro(s) recebido(s) para o relatorio de saldo devedor presente.", len(filtros))

    for indice, filtro in enumerate(filtros, start=1):
        resumo = (
            f"[saldo_devedor][diagnostico][filtro={indice}] "
            f"id={filtro.get('id') or filtro.get('id_filtro')} | "
            f"nome={filtro.get('nome') or filtro.get('nome_filtro')} | "
            f"campo={filtro.get('campo_identificador')} | "
            f"acao={filtro.get('acao')} | "
            f"valor={filtro.get('valor')}"
        )
        print(resumo, flush=True)
        logger.info(resumo)


def _consolidar_planilhas(
    arquivos,
    relatorio_config,
    empresa_id,
    download_dir,
    sanitizar_nome_arquivo,
    montar_caminho_download_unico,
    logger,
):
    workbook_destino = Workbook()
    worksheet_destino = workbook_destino.active or workbook_destino.create_sheet()
    worksheet_destino.title = "Consolidado"

    escreveu_linhas = False

    for indice, arquivo in enumerate(arquivos, start=1):
        logger.info("Consolidando arquivo %s/%s: %s", indice, len(arquivos), arquivo)
        linhas = list(_iterar_linhas_planilha(arquivo))
        if linhas:
            linhas = linhas[:-1]
            logger.info("Linha totalizadora removida do arquivo antes da consolidacao: %s", arquivo)
        for numero_linha, linha in enumerate(linhas, start=1):
            if indice > 1 and numero_linha == 1:
                continue
            worksheet_destino.append(list(linha))
            escreveu_linhas = True

    if not escreveu_linhas:
        logger.warning("Nenhuma linha foi consolidada para o saldo devedor presente.")
        return None

    nome_relatorio = sanitizar_nome_arquivo(relatorio_config.get("nome") or "relatorio")
    empresa = sanitizar_nome_arquivo(str(empresa_id or "sem_empresa"))
    nome_base = f"{nome_relatorio}_{empresa}_consolidado"
    destino = montar_caminho_download_unico(Path(download_dir), nome_base, ".xlsx")
    workbook_destino.save(destino)
    logger.info("Arquivo consolidado salvo em %s.", destino)
    return destino


def _remover_arquivos_intermediarios(arquivos, consolidado, logger):
    caminho_consolidado = Path(consolidado).resolve() if consolidado else None

    for arquivo in arquivos:
        caminho = Path(arquivo)
        try:
            if caminho_consolidado and caminho.resolve() == caminho_consolidado:
                continue
            if caminho.exists():
                caminho.unlink()
                logger.info("Arquivo intermediario removido apos consolidacao: %s", caminho)
        except OSError:
            logger.exception("Nao foi possivel remover o arquivo intermediario %s.", caminho)


def _iterar_linhas_planilha(arquivo):
    sufixo = arquivo.suffix.lower()
    if sufixo == ".csv":
        with arquivo.open("r", encoding="utf-8-sig", newline="") as csv_file:
            for linha in csv_reader(csv_file):
                yield linha
        return

    workbook = load_workbook(arquivo, data_only=True)
    # active pode ser None (ex.: aba de grafico selecionada); cai na primeira planilha.
    worksheet = workbook.active or workbook.worksheets[0]
    for linha in worksheet.iter_rows(values_only=True):
        yield linha
